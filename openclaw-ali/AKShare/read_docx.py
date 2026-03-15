# -*- coding: utf-8 -*-
from docx import Document
import sys

if sys.platform == 'win32':
    import io
    sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

doc_path = r'C:\AI-Agent-Local\openclaw-ali\AKShare\AKShare股票数据全面接入服务器使用手册.docx'
doc = Document(doc_path)

print('=' * 100)
print('  AKShare股票数据全面接入服务器使用手册')
print('=' * 100)
print()

for i, para in enumerate(doc.paragraphs):
    text = para.text.strip()
    if text:
        print(text)

# 读取表格
if doc.tables:
    print('\n')
    print('=' * 100)
    print('  表格内容')
    print('=' * 100)
    print()

    for table_idx, table in enumerate(doc.tables):
        print(f'--- 表格 {table_idx + 1} ---')
        print()

        for row in table.rows:
            row_data = []
            for cell in row.cells:
                row_data.append(cell.text.strip())
            print(' | '.join(row_data))
        print()
