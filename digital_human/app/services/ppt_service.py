# -*- coding: utf-8 -*-
"""PPT 出片服务 (2026-08-20 v2) — 按 shape 原始坐标/字号/颜色/背景还原页面.

v2 (还原布局): 不再简陋重排文字, 而是用 python-pptx 读取每个 shape 的
  位置/大小/字号/颜色 + 图片 + 页面背景图, 按原始坐标在 1920x1080 HTML 里
  精准摆放 → Chrome headless 截图 → ffmpeg mp4 (Ken Burns + 淡入).
"""
from __future__ import annotations

import base64
import logging
import re
import subprocess
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from app.config import get_config

logger = logging.getLogger(__name__)

__all__ = ["parse_pptx", "build_slide_html", "render_slide_mp4", "clean_notes"]

_CHROME = r"C:/Program Files/Google/Chrome/Application/chrome.exe"
_EMU_PER_PX = 9525  # 12192000 EMU / 1280 逻辑px → 按 1920 宽等比: 1920*9525=18288000? 用比例
# 幻灯片 12192000x6858000 EMU, 目标 1920x1080: 比例系数
_PX_PER_EMU_W = 1920 / 12192000
_PX_PER_EMU_H = 1080 / 6858000

# 备注脏前缀/尾标
_NOTE_PREFIX_RE = re.compile(r"^\s*\d+\.\d+\.\d+\s*")
_NOTE_TAIL_RE = re.compile(r"\s*‹#›\s*$")


@dataclass
class TextBlock:
    text: str
    left: float
    top: float
    width: float
    height: float
    font_size_pt: float  # None → 默认
    color: str | None  # hex
    bold: bool = False
    shape_id: int = 0  # parse 时按 shape 迭代序赋 (动画 targeting)


@dataclass
class ShapeBlock:
    """装饰自选图形 (卡片底/暗化层等): 填充色+alpha, 圆角, 位置. 不进动画层 (基底)."""
    left: float
    top: float
    width: float
    height: float
    fill_hex: str | None  # e.g. '000000' (无#)
    fill_alpha: float = 1.0  # 0-1 (XML alpha/100000)
    rounded: bool = False
    shape_id: int = 0


@dataclass
class ImageBlock:
    b64_data: str
    left: float
    top: float
    width: float
    height: float
    shape_id: int = 0
    role: str = "content"  # content(搜图) | style(生图), v1 默认 content


@dataclass
class Slide:
    index: int
    texts: list[str] = field(default_factory=list)  # 纯文本列表 (预览用)
    text_blocks: list[TextBlock] = field(default_factory=list)
    image_blocks: list[ImageBlock] = field(default_factory=list)
    shape_blocks: list[ShapeBlock] = field(default_factory=list)  # 装饰自选图形
    background_b64: str | None = None  # 页面背景图
    notes: str = ""


def clean_notes(text: str) -> str:
    """清洗备注台词: 去豆包时间戳前缀 + ‹#› 页码标记."""
    t = _NOTE_TAIL_RE.sub("", text or "")
    t = _NOTE_PREFIX_RE.sub("", t)
    return t.strip()


def _pt_to_px(pt: float) -> float:
    """pt (1/72 英寸) → px. 画布 1920px = 13.33in 幻灯片 → 144dpi, 故 ×2.

    2026-08-21 修复: 原 96dpi(×1.333)导致字号缩到 67%, 文字偏小+填不满框=排布错位.
    """
    return max(pt, _MIN_FONT_PT) * 144 / 72  # == pt * 2 (6pt 下限)


# 最小可读字号 (2026-08-21): <6pt 在手机屏无法看, 作为小字底线
_MIN_FONT_PT = 6.0


def _color_hex(rgb) -> str | None:
    try:
        if rgb is None:
            return None
        return str(rgb)
    except Exception:
        return None


def _bg_image(slide, prs) -> str | None:
    """提取页面背景图 (p:cSld/p:bg/p:bgPr/a:blipFill/a:blip → rel → base64).

    精确锁定 p:bg 路径, 杜绝把 shape 配图当背景 (旧宽松正则 bug).
    p:bgRef (引母版) 无独立背景图 → 返回 None (由 skin 包统一补).
    """
    from pptx.oxml.ns import qn

    try:
        sld = slide._element
        cSld = sld.find(qn("p:cSld"))
        if cSld is None:
            return None
        bg = cSld.find(qn("p:bg"))
        if bg is None:
            return None
        bg_pr = bg.find(qn("p:bgPr"))
        if bg_pr is None:
            return None  # p:bgRef 引母版, 无独立图
        blip_fill = bg_pr.find(qn("a:blipFill"))
        blip = blip_fill.find(qn("a:blip")) if blip_fill is not None else None
        if blip is None:
            return None
        rid = blip.get(qn("r:embed"))
        if not rid:
            return None
        target_part = slide.part.related_part(rid)
        blob = target_part.blob
        ctype = target_part.content_type
        return f"data:{ctype};base64,{base64.b64encode(blob).decode()}"
    except Exception as exc:
        logger.warning("[ppt] 背景图提取失败: %s", exc)
        return None


def _shape_block(shape, sid: int) -> ShapeBlock | None:
    """自选图形 → 装饰块 (卡片底/暗化层/细线装饰). 填充/线色读 XML (含 alpha)."""
    from pptx.oxml.ns import qn
    try:
        sp_pr = shape._element.find(qn("p:spPr"))
        if sp_pr is None:
            return None
        fill = sp_pr.find(qn("a:solidFill"))
        hexval = None
        alpha = 1.0
        if fill is not None:
            clr = fill.find(qn("a:srgbClr"))
            if clr is not None:
                hexval = clr.get("val")
                a = clr.find(qn("a:alpha"))
                if a is not None and a.get("val"):
                    alpha = int(a.get("val")) / 100000.0
            elif shape.fill.type is not None:
                try:
                    hexval = str(shape.fill.fore_color.rgb)
                except Exception:
                    hexval = None
        else:
            # 无填充 → 有可见线(a:ln/a:solidFill)也算装饰块 (细线/边框, PPT 常见)
            ln = sp_pr.find(qn("a:ln"))
            if ln is not None:
                lf = ln.find(qn("a:solidFill"))
                if lf is not None:
                    clr = lf.find(qn("a:srgbClr"))
                    if clr is not None:
                        hexval = clr.get("val")
                        a = clr.find(qn("a:alpha"))
                        if a is not None and a.get("val"):
                            alpha = int(a.get("val")) / 100000.0
        if not hexval:
            return None
        rounded = False
        geom = sp_pr.find(qn("a:prstGeom"))
        if geom is not None:
            prst = geom.get("prst") or ""
            # roundRect 半径由 adj 控制; adj=0 → 直角 (本 deck 全为 0)
            if "round" in prst.lower():
                adj = 0.0
                av_lst = geom.find(qn("a:avLst"))
                if av_lst is not None:
                    for gd in av_lst.findall(qn("a:gd")):
                        if gd.get("name") == "adj":
                            fmla = (gd.get("fmla") or "").split()
                            if len(fmla) == 3:
                                try:
                                    adj = float(fmla[2])
                                except ValueError:
                                    adj = 0.0
                rounded = adj > 0
        return ShapeBlock(
            left=shape.left, top=shape.top,
            width=shape.width, height=shape.height,
            fill_hex=hexval, fill_alpha=alpha, rounded=rounded, shape_id=sid,
        )
    except Exception:
        return None


def parse_pptx(path: str | Path) -> list[Slide]:
    """解析 pptx → 每页结构化 shape (坐标/字号/颜色/图/背景)."""
    from pptx import Presentation

    prs = Presentation(str(path))
    slides: list[Slide] = []
    for i, slide in enumerate(prs.slides, 1):
        text_blocks: list[TextBlock] = []
        image_blocks: list[ImageBlock] = []
        shape_blocks: list[ShapeBlock] = []
        texts: list[str] = []
        for sid, shape in enumerate(slide.shapes):
            try:
                if shape.has_text_frame and shape.text_frame.text.strip():
                    txt = shape.text_frame.text.strip()
                    texts.append(txt)
                    # 取首段首 run 的字号/颜色
                    sz_pt = None
                    color = None
                    bold = False
                    for para in shape.text_frame.paragraphs:
                        for run in para.runs:
                            if run.font.size:
                                sz_pt = run.font.size.pt
                            color = _color_hex(run.font.color.rgb) if (run.font.color and run.font.color.rgb) else color
                            bold = bool(run.font.bold)
                            if sz_pt:
                                break
                        if sz_pt:
                            break
                    tb = TextBlock(
                        text=txt,
                        left=shape.left, top=shape.top,
                        width=shape.width, height=shape.height,
                        font_size_pt=sz_pt, color=color, bold=bold,
                        shape_id=sid,
                    )
                    text_blocks.append(tb)
                if shape.shape_type == 13:  # PICTURE
                    try:
                        img = shape.image
                        b64 = f"data:{img.content_type};base64,{base64.b64encode(img.blob).decode()}"
                        image_blocks.append(ImageBlock(
                            b64_data=b64, left=shape.left, top=shape.top,
                            width=shape.width, height=shape.height,
                            shape_id=sid,
                        ))
                    except Exception as exc:
                        logger.warning("[ppt] 页%d 图提取失败: %s", i, exc)
                # 装饰自选图形 (卡片底/暗化层): 有 solidFill 即捕获, 色值读 XML (含 alpha)
                sb = _shape_block(shape, sid)
                if sb:
                    shape_blocks.append(sb)
            except Exception as exc:
                logger.warning("[ppt] 页%d shape 解析失败: %s", i, exc)
        notes = ""
        try:
            if slide.has_notes_slide:
                notes = clean_notes(slide.notes_slide.notes_text_frame.text)
        except Exception as exc:
            logger.warning("[ppt] 页%d 备注失败: %s", i, exc)
        bg = _bg_image(slide, prs)
        if bg is None:
            # 无 p:bg → 全幅 picture shape 即背景 (PPT 常见做法: 铺满整页的图不是 p:bg)
            # 全幅图升为背景, 并从前景 image_blocks 移除 (避免元素层重复渲染)
            for ib in list(image_blocks):
                if (ib.width >= 0.9 * prs.slide_width
                        and ib.height >= 0.9 * prs.slide_height):
                    bg = ib.b64_data
                    image_blocks.remove(ib)
                    break
        slides.append(Slide(
            index=i, texts=texts, text_blocks=text_blocks,
            image_blocks=image_blocks, shape_blocks=shape_blocks,
            background_b64=bg, notes=notes,
        ))
    return slides


def _esc(s: str) -> str:
    return s.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def _css_color(c) -> str:
    """确保 CSS 色值带 '#' 前缀 (parse 给 'FFFFFF' 不带 # → 非法 CSS → 浏览器回落黑字)."""
    if not c:
        return "#3d2b1f"
    s = str(c)
    return s if s.startswith("#") else "#" + s


def _px_left(v): return round(v * _PX_PER_EMU_W)
def _px_top(v): return round(v * _PX_PER_EMU_H)
def _px_w(v): return round(v * _PX_PER_EMU_W)
def _px_h(v): return round(v * _PX_PER_EMU_H)


def build_slide_html(slide: Slide, *, skin=None, width: int = 1920, height: int = 1080) -> str:
    """按 shape 原始坐标/字号/颜色/背景生成带分层入场动画的 HTML.

    skin (SeriesSkinPack|None): 非 None 时用 skin 的 font_family/animation_profile
    (色/字号/背景已在 apply_skin_to_slides 阶段覆盖到 slide 内).
    动画: shape 迭代序= back-to-front 堆叠, 按 idx*stagger 错峰入场.
    背景层恒显 (无 anim-layer) 保证首帧完整. 注入 window.__hf.seek(t) JS hook
    供 Playwright seek-and-snap 逐帧定位.
    """
    def px_left(v): return round(v * _PX_PER_EMU_W)
    def px_top(v): return round(v * _PX_PER_EMU_H)
    def px_w(v): return round(v * _PX_PER_EMU_W)
    def px_h(v): return round(v * _PX_PER_EMU_H)

    font_family = skin.font_family if skin else (
        'system-ui,"PingFang SC","Microsoft YaHei",sans-serif')
    anim = skin.animation_profile if skin else {"stagger": 0.22, "in_dur": 0.55}
    stagger = float(anim.get("stagger", 0.22))

    # 背景 (图片或纯色)
    if slide.background_b64:
        bg_style = (f"background-image:url('{slide.background_b64}');"
                    f"background-size:{width}px {height}px;background-position:center;")
    elif skin and skin.color_tokens:
        bg_style = f"background:{skin.color_tokens[-1]};"  # 末位色做背景底色
    else:
        bg_style = "background:#f5f0e6;"

    # 统一 shapes 迭代序 (text+image 按 shape_id 混排保 z-order)
    blocks = []
    for tb in slide.text_blocks:
        blocks.append(("text", tb))
    for ib in slide.image_blocks:
        blocks.append(("img", ib))
    blocks.sort(key=lambda b: b[1].shape_id)

    els = []
    for idx, (kind, b) in enumerate(blocks):
        delay = round(idx * stagger, 3)
        # 首块用纯 fade (避免最底层位移突兀), 余块 fade-slide-up
        anim_name = "hf-fade-in" if idx == 0 else "hf-fade-slide-up"
        anim_style = (f"animation-name:{anim_name};animation-delay:{delay}s;")
        if kind == "text":
            fs = _pt_to_px(b.font_size_pt) if b.font_size_pt else 28
            color = _css_color(b.color or (skin.color_tokens[0] if skin and skin.color_tokens else None))
            els.append(
                f'<div id="anim-{b.shape_id}" class="anim-layer" '
                f'style="position:absolute;left:{px_left(b.left)}px;top:{px_top(b.top)}px;'
                f'width:{px_w(b.width)}px;height:{px_h(b.height)}px;'
                f'font-size:{fs}px;line-height:1.3;color:{color};'
                f'{"font-weight:700;" if b.bold else ""}'
                f'overflow:hidden;word-wrap:break-word;white-space:pre-wrap;box-sizing:border-box;'
                f'{anim_style}">{_esc(b.text)}</div>'
            )
        else:
            els.append(
                f'<img id="anim-{b.shape_id}" class="anim-layer" '
                f'style="position:absolute;left:{px_left(b.left)}px;top:{px_top(b.top)}px;'
                f'width:{px_w(b.width)}px;height:{px_h(b.height)}px;'
                f'object-fit:contain;{anim_style}" src="{b.b64_data}">'
            )

    return f"""<!DOCTYPE html>
<html><head><meta charset="utf-8"><style>
  html,body{{margin:0;padding:0;width:{width}px;height:{height}px;overflow:hidden;
    font-family:{font_family};}}
  .anim-layer{{opacity:0;animation-fill-mode:both;animation-duration:0.55s;
    animation-timing-function:cubic-bezier(0.22,1,0.36,1);}}
  @keyframes hf-fade-slide-up{{from{{opacity:0;transform:translateY(24px)}}to{{opacity:1;transform:none}}}}
  @keyframes hf-fade-in{{from{{opacity:0}}to{{opacity:1}}}}
</style></head><body>
  <div id="stage" style="position:relative;width:{width}px;height:{height}px;{bg_style}">
    {''.join(els)}
  </div>
<script>window.__hf={{seek:function(t){{document.getAnimations().forEach(function(a){{try{{a.currentTime=t*1000;a.pause()}}catch(e){{}}}})}}}};</script>
</body></html>"""


def _page_shell(width: int, height: int, bg_style: str = "background:transparent") -> str:
    return (f"<!DOCTYPE html><html><head><meta charset='utf-8'><style>"
            f"html,body{{margin:0;padding:0;width:{width}px;height:{height}px;"
            f"overflow:hidden;{bg_style}}}</style></head><body>")


def build_base_html(slide: Slide, width: int = 1920, height: int = 1080) -> str:
    """基底层 HTML (2026-08-21 元素级拆解): 背景图 + 装饰自选图形合成.

    不透明白底 (底层). 文字/前景图作为独立透明层叠其上.
    """
    els = []
    # 背景图铺满
    if slide.background_b64:
        els.append(f"<img style='position:absolute;left:0;top:0;width:{width}px;"
                   f"height:{height}px;object-fit:fill' src='{slide.background_b64}'>")
    else:
        els.append(f"<div style='position:absolute;left:0;top:0;width:{width}px;"
                   f"height:{height}px;background:#f5f0e6'></div>")
    # 装饰自选图形 (含全幅暗化层/卡片底), 按 shape_id 序叠放
    for sb in slide.shape_blocks:
        alpha = max(0.0, min(1.0, sb.fill_alpha))
        r = f"border-radius:{min(_px_w(sb.width), _px_h(sb.height)) * 0.12}px" if sb.rounded else ""
        els.append(
            f"<div style='position:absolute;left:{_px_left(sb.left)}px;top:{_px_top(sb.top)}px;"
            f"width:{_px_w(sb.width)}px;height:{_px_h(sb.height)}px;"
            f"background:{_css_color(sb.fill_hex)};opacity:{alpha:.3f};{r}'></div>")
    return _page_shell(width, height) + "".join(els) + "</body></html>"


def build_text_element_html(tb: TextBlock, width: int = 1920, height: int = 1080) -> str:
    """单文字块透明层 HTML: 块在自身坐标渲染 (像素级), 供剪映叠层逐级入场."""
    fs = _pt_to_px(tb.font_size_pt) if tb.font_size_pt else 28
    div = (
        f"<div style='position:absolute;left:{_px_left(tb.left)}px;top:{_px_top(tb.top)}px;"
        f"width:{_px_w(tb.width)}px;height:{_px_h(tb.height)}px;"
        f"font-size:{fs}px;line-height:1.3;color:{_css_color(tb.color)};"
        f"{'font-weight:700;' if tb.bold else ''}"
        f"font-family:'Noto Sans SC','Source Han Sans SC','Microsoft YaHei','PingFang SC',sans-serif;"
        f"overflow:hidden;word-wrap:break-word;white-space:pre-wrap;box-sizing:border-box'>"
        f"{_esc(tb.text)}</div>"
    )
    return _page_shell(width, height) + div + "</body></html>"


def build_image_element_html(ib: ImageBlock, width: int = 1920, height: int = 1080) -> str:
    """单前景图透明层 HTML: 图在自身坐标渲染 (object-fit:contain 保形)."""
    img = (
        f"<img style='position:absolute;left:{_px_left(ib.left)}px;top:{_px_top(ib.top)}px;"
        f"width:{_px_w(ib.width)}px;height:{_px_h(ib.height)}px;"
        f"object-fit:contain' src='{ib.b64_data}'>"
    )
    return _page_shell(width, height) + img + "</body></html>"


def render_slide_png(slide_html: str, work_dir: Path, output_png: Path,
                     width: int = 1920, height: int = 1080) -> Path:
    """渲染单页为 1 张最终态静态 PNG (2026-08-21).

    替代逐帧捕获: Playwright seek 到动画终态截图 (~1s/页). 动画不烧进帧里,
    由剪映 IntroType 入场动画 (jy 模式) 或 ffmpeg zoompan (auto 模式) 提供.
    失败回退 Chrome headless --virtual-time-budget 截图.
    """
    work_dir.mkdir(parents=True, exist_ok=True)
    html_path = work_dir / "slide.html"
    html_path.write_text(slide_html, encoding="utf-8")
    try:
        from app.services.ppt_frame_capture import capture_slide_png
        return capture_slide_png(html_path, output_png, width=width, height=height)
    except Exception as exc:
        logger.warning("[ppt] Playwright 静态帧失败, 回退 Chrome 截图: %s", exc)
        return _chrome_screenshot_png(slide_html, work_dir, output_png, width, height)


def png_to_zoompan_mp4(png_path: Path, output_mp4: Path, duration_sec: float,
                       width: int = 1920, height: int = 1080) -> Path:
    """静态 PNG → Ken Burns 缩放+淡入 mp4 (auto 模式动画层)."""
    from app.infrastructure.ffmpeg import run_ffmpeg
    dur = max(1.5, duration_sec)
    fps = 25
    frames = int(dur * fps)
    vf = (
        f"zoompan=z='min(zoom+0.0004,1.08)':d={frames}:x='iw/2-(iw/zoom/2)':"
        f"y='ih/2-(ih/zoom/2)':s={width}x{height}:fps={fps},"
        f"fade=t=in:st=0:d=0.4,fade=t=out:st={max(0, dur-0.4):.2f}:d=0.4"
    )
    run_ffmpeg([
        "ffmpeg", "-y", "-loglevel", "error",
        "-loop", "1", "-i", str(png_path),
        "-vf", vf, "-t", f"{dur:.2f}",
        "-c:v", "libx264", "-pix_fmt", "yuv420p", "-r", str(fps),
        str(output_mp4),
    ], timeout=120)
    return output_mp4


def _chrome_screenshot_png(slide_html: str, work_dir: Path, output_png: Path,
                           width: int, height: int) -> Path:
    """Chrome headless 截最终态 PNG (动画走完: --virtual-time-budget 拉满)."""
    from app.infrastructure.ffmpeg import run_ffmpeg
    html_path = work_dir / "slide.html"
    html_path.write_text(slide_html, encoding="utf-8")
    cmd = [
        _CHROME, "--headless=new", "--disable-gpu", "--no-sandbox",
        f"--screenshot={output_png}", f"--window-size={width},{height}",
        "--virtual-time-budget=15000", html_path.as_uri(),
    ]
    try:
        run_ffmpeg(cmd, timeout=60)
    except RuntimeError:
        pass  # Chrome 截图失败时直接判文件
    if not output_png.exists():
        raise RuntimeError("Chrome 截图失败 (static fallback)")
    return output_png


def render_slide_mp4(slide_html: str, work_dir: Path, output_mp4: Path,
                     duration_sec: float, width: int = 1920, height: int = 1080) -> Path:
    """渲染单页 mp4 — 主路径 Playwright seek-and-snap 逐帧 (分层入场动效).

    失败回退 _render_slide_mp4_static (旧 Chrome 截图 + zoompan Ken Burns),
    不阻断产线 (风险点 3: 生产若未装 playwright chromium).
    """
    work_dir.mkdir(parents=True, exist_ok=True)
    html_path = work_dir / "slide.html"
    html_path.write_text(slide_html, encoding="utf-8")

    dur = max(1.5, duration_sec)
    try:
        from app.services.ppt_frame_capture import capture_slide_frames
        from app.infrastructure.ffmpeg import frames_to_mp4
        frames_dir = work_dir / "frames"
        if frames_dir.exists():
            for f in frames_dir.iterdir():
                f.unlink()
        capture_slide_frames(html_path, frames_dir, dur, fps=25, width=width, height=height)
        frames_to_mp4(frames_dir, "frame_%06d.png", output_mp4, fps=25,
                      width=width, height=height)
        return output_mp4
    except Exception as exc:
        logger.warning("[ppt] Playwright 动画渲染失败, 回退 static zoompan: %s", exc)
        return _render_slide_mp4_static(slide_html, work_dir, output_mp4, dur, width, height)


def _render_slide_mp4_static(slide_html: str, work_dir: Path, output_mp4: Path,
                             duration_sec: float, width: int = 1920, height: int = 1080) -> Path:
    """旧路径 fallback: Chrome 单 PNG → zoompan Ken Burns + 淡入."""
    png = _chrome_screenshot_png(slide_html, work_dir, work_dir / "slide.png", width, height)
    return png_to_zoompan_mp4(png, output_mp4, duration_sec, width, height)
